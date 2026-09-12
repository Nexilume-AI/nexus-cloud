"""Actual owner HTTP, encrypted import queue and loopback model discovery."""
import io
import json
from datetime import timedelta
from uuid import uuid4
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.test import APIClient
from apps.audit.models import AuditLog
from apps.common.crypto import decrypt_secret
from apps.providers.models import ProviderAccount, ProviderImportBatch, ProviderRuntimeAccount, ProviderRuntimeModelOffer
from apps.providers.tasks import clear_expired_provider_imports
from .provider_http_fixture import ProviderHTTPFixture

ROOT = '/api/v1/provider-connections/imports/'
HEAD = 'api_ref,name,base_url,api_key\n'


@override_settings(ROOT_URLCONF='nexus_personal.urls', NEXUS_PROVIDER_ALLOW_HTTP=True,
                   NEXUS_PROVIDER_ALLOWED_PRIVATE_HOSTS='127.0.0.1')
class PersonalProviderImportHTTPTests(ProviderHTTPFixture, TestCase):
    def setUp(self):
        super().setUp()
        self.client = APIClient(enforce_csrf_checks=True)
        self.client.force_login(self.installation.owner)
        self.client.get('/api/v1/public/bootstrap/')
        self.headers = {'HTTP_X_CSRFTOKEN': self.client.cookies['csrftoken'].value, 'HTTP_ORIGIN': 'http://testserver'}

    def table(self, count=1, key='local-provider-test-key'):
        return HEAD + ''.join(f'import-{i},Imported {i},http://127.0.0.1:{self.upstream.server_port}/v1,{key}\n' for i in range(count))

    def post(self, path, body=None):
        return self.client.post(path, body or {}, format='json', **self.headers)

    def preview(self, table=None, **kwargs):
        response = self.post(ROOT + 'preview/', {'apis_text': table or self.table(), 'request_key': str(uuid4()), **kwargs})
        self.assertEqual(response.status_code, 201, response.data)
        self.assert_safe(response.data)
        return response.data

    def commit(self, batch, **kwargs):
        response = self.post(ROOT + batch['id'] + '/commit/', kwargs)
        self.assertEqual(response.status_code, 200, response.data)
        self.assert_safe(response.data)
        return response.data

    def assert_safe(self, value):
        text = json.dumps(value, default=str)
        for secret in ('local-provider-test-key', 'encrypted_payload', 'encrypted_key', 'Authorization'):
            self.assertNotIn(secret, text)

    def test_capabilities_templates_real_batch_chunks_and_post_import_discovery(self):
        capabilities = self.client.get(ROOT + 'capabilities/')
        self.assertEqual(capabilities.status_code, 200, capabilities.data)
        self.assertTrue(capabilities.data['can_import'])
        self.assertTrue(capabilities.data['can_update_existing'])
        self.assertEqual(capabilities['Cache-Control'], 'no-store')
        preview = self.preview(self.table(6))
        record = ProviderImportBatch.objects.get(pk=preview['id'])
        self.assertNotIn('local-provider-test-key', record.encrypted_payload)
        self.assertIn('local-provider-test-key', decrypt_secret(record.encrypted_payload))
        self.assertEqual(ProviderAccount.objects.count(), 0)
        first = self.commit(preview)
        self.assertEqual(first['remaining'], 1)
        self.assertEqual(ProviderAccount.objects.count(), 5)
        second = self.commit(preview)
        self.assertEqual(second['status'], 'complete')
        self.assertEqual(ProviderAccount.objects.count(), 6)
        self.assertEqual(self.commit(preview), second)
        record.refresh_from_db()
        self.assertEqual(record.encrypted_payload, '')
        self.assertEqual(self.calls, [])
        self.assertEqual(ProviderRuntimeModelOffer.objects.count(), 0)
        account_id = second['results'][0]['connection_id']
        started = self.post(f'/api/v1/provider-connections/{account_id}/start/')
        self.assertEqual(started.status_code, 200, started.data)
        self.assertEqual(started.data['models'][0]['upstream_model_id'], 'personal-model')
        self.assertIn(('GET', '/v1/models'), self.calls)
        report = self.client.get(ROOT + preview['id'] + '/report/')
        self.assertEqual(report.status_code, 200)
        self.assertNotIn(b'local-provider-test-key', report.content)
        for event in AuditLog.objects.filter(action__startswith='providers.import.'):
            self.assert_safe(event.metadata)

    def test_idempotent_preview_conflicts_and_explicit_partial_confirmation(self):
        key = str(uuid4())
        body = {'apis_text': self.table(), 'request_key': key}
        first = self.post(ROOT + 'preview/', body)
        again = self.post(ROOT + 'preview/', body)
        self.assertEqual(first.data['id'], again.data['id'])
        self.assertEqual(ProviderImportBatch.objects.count(), 1)
        self.assertEqual(self.post(ROOT + 'preview/', {**body, 'apis_text': self.table(2)}).status_code, 400)
        mixed = self.preview(self.table() + 'bad,Invalid,ftp://invalid.test/v1,local-provider-test-key\n')
        self.assertEqual(mixed['invalid'], 1)
        self.assertEqual(self.post(ROOT + mixed['id'] + '/commit/').status_code, 400)
        self.assertEqual(ProviderAccount.objects.count(), 0)
        result = self.commit(mixed, allow_partial=True)
        self.assertEqual([item['status'] for item in result['results']], ['created', 'invalid'])

    def test_duplicate_skip_update_confirmation_empty_key_and_stale_revision(self):
        first = self.commit(self.preview())
        account = ProviderAccount.objects.get(pk=first['results'][0]['connection_id'])
        self.assertEqual(self.preview()['results'][0]['status'], 'skipped')
        update = self.preview(self.table(key='').replace('Imported 0', 'Renamed'), duplicate_mode='update')
        self.assertEqual(update['results'][0]['action'], 'update')
        self.assertEqual(self.post(ROOT + update['id'] + '/commit/').status_code, 400)
        self.assertEqual(self.commit(update, confirm_updates=True)['results'][0]['status'], 'updated')
        account.refresh_from_db()
        self.assertEqual(account.name, 'Renamed')
        self.assertEqual(decrypt_secret(account.encrypted_key), 'local-provider-test-key')
        stale = self.preview(self.table(key=''), duplicate_mode='update')
        account.name = 'External edit'
        account.save(update_fields=['name', 'updated_at'])
        self.assertEqual(self.commit(stale, confirm_updates=True)['results'][0]['status'], 'failed')
        account.refresh_from_db()
        self.assertEqual(account.name, 'External edit')

    def test_current_owner_and_context_are_rechecked_at_commit(self):
        first = self.commit(self.preview())
        account = ProviderAccount.objects.get(pk=first['results'][0]['connection_id'])
        batch = self.preview(self.table(key=''), duplicate_mode='update')
        # Simulate a changed runtime owner without bumping the preview revision.
        ProviderRuntimeAccount.objects.filter(source_provider_account=account).update(owner=self.other)
        self.assertEqual(self.commit(batch, confirm_updates=True)['results'][0]['status'], 'failed')
        self.assertEqual(self.preview()['results'][0]['status'], 'invalid')
        self.assertEqual(account.name, 'Imported 0')
        unknown = APIClient()
        unknown.force_login(self.other)
        self.assertEqual(unknown.get(ROOT + batch['id'] + '/').status_code, 401)
        self.assertEqual(self.client.get(ROOT + batch['id'] + '/', HTTP_X_NEXUS_PROJECT='foreign').status_code, 403)
        self.assertEqual(self.client.get(ROOT + str(uuid4()) + '/').status_code, 404)

    def test_csrf_authentication_and_file_parsing_bounds(self):
        self.assertEqual(APIClient().get(ROOT + 'capabilities/').status_code, 401)
        self.assertEqual(self.client.post(ROOT + 'preview/', {}, format='json').status_code, 403)
        for extra in ({'models': ['forged']}, {'ownership': {'scope': 'organization'}}, {'tenant_id': str(uuid4())}):
            self.assertEqual(self.post(ROOT + 'preview/', {'apis_text': self.table(), 'request_key': str(uuid4()), **extra}).status_code, 400)
        template = self.client.get(ROOT + 'template/')
        self.assertEqual(template.status_code, 200)
        book = load_workbook(io.BytesIO(template.content))
        self.assertEqual(book.sheetnames, ['APIs', 'Guide'])
        book['APIs'].append(['xlsx-import', 'XLSX', f'http://127.0.0.1:{self.upstream.server_port}/v1', 'local-provider-test-key'])
        stream = io.BytesIO()
        book.save(stream)
        valid = self.client.post(ROOT + 'preview/', {'request_key': str(uuid4()),
            'file': SimpleUploadedFile('apis.xlsx', stream.getvalue())}, format='multipart', **self.headers)
        self.assertEqual(valid.status_code, 201, valid.data)
        self.assert_safe(valid.data)
        book['APIs']['B2'] = '=1+1'
        stream = io.BytesIO()
        book.save(stream)
        book.close()
        for upload in (SimpleUploadedFile('apis.xlsx', stream.getvalue()), SimpleUploadedFile('broken.xlsx', b'bad'),
                       SimpleUploadedFile('big.csv', b'x' * (2_097_152 + 1))):
            rejected = self.client.post(ROOT + 'preview/', {'request_key': str(uuid4()), 'file': upload}, format='multipart', **self.headers)
            self.assertEqual(rejected.status_code, 400, rejected.data)
            self.assert_safe(rejected.data)
        self.assertEqual(self.post(ROOT + 'preview/', {'apis_text': self.table(101), 'request_key': str(uuid4())}).status_code, 400)
        self.assertEqual(self.calls, [])

    def test_expiry_discard_and_real_periodic_cleanup_erase_only_payloads(self):
        batch = self.preview()
        path = ROOT + batch['id'] + '/'
        self.assertEqual(self.client.delete(path, **self.headers).status_code, 200)
        self.assertEqual(ProviderImportBatch.objects.get(pk=batch['id']).encrypted_payload, '')
        self.assertEqual(self.post(path + 'commit/').status_code, 400)
        expiry = self.preview()
        ProviderImportBatch.objects.filter(pk=expiry['id']).update(expires_at=timezone.now() - timedelta(seconds=1))
        self.assertEqual(clear_expired_provider_imports(), {'cleared': 1})
        self.assertEqual(clear_expired_provider_imports(), {'cleared': 0})
        self.assertEqual(ProviderImportBatch.objects.get(pk=expiry['id']).encrypted_payload, '')
        self.assertEqual(self.client.get(ROOT + expiry['id'] + '/').data['status'], 'expired')
        self.assertEqual(len(self.client.get(ROOT).data), 2)
        self.assertEqual(ProviderAccount.objects.count(), 0)

    def test_capacity_failure_retains_retry_payload_and_cannot_duplicate_success(self):
        batch = self.preview(self.table(2))
        with override_settings(NEXUS_PERSONAL_MODEL_LIMITS={'models.provider_connections': 1}):
            first = self.commit(batch)
        self.assertEqual([row['status'] for row in first['results']], ['created', 'failed'])
        self.assertNotIn('Plan', str(first))
        self.assertEqual(ProviderAccount.objects.count(), 1)
        finished = self.commit(batch, retry_failed=True)
        self.assertEqual([row['status'] for row in finished['results']], ['created', 'created'])
        self.assertEqual(ProviderAccount.objects.count(), 2)
        self.assertEqual(ProviderImportBatch.objects.get(pk=batch['id']).encrypted_payload, '')
