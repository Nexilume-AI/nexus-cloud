"""Bounded spreadsheet parsing. Never evaluate formulas, fetch links or save uploads."""
from __future__ import annotations

import csv
import io
import re
from xml.etree import ElementTree
from zipfile import BadZipFile, ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from rest_framework.exceptions import ValidationError

API_COLUMNS = ["api_ref", "name", "base_url", "api_key"]
MAX_BYTES = 2 * 1024 * 1024
MAX_EXPANDED = 8 * 1024 * 1024
MAX_APIS = 100


def _rows(values, columns, limit, label):
    iterator = iter(values)
    header = next(iterator, [])
    names = [str(value or "").strip().lstrip("\ufeff") for value in header]
    while names and not names[-1]:
        names.pop()
    if not names or len(names) != len(set(names)) or any(name not in columns for name in names):
        raise ValidationError({label: "Use the template headers, without duplicate or unknown columns."})
    required = API_COLUMNS[:3]
    if not set(required).issubset(names):
        raise ValidationError({label: "Required template columns are missing."})
    result = []
    for line, cells in enumerate(iterator, 2):
        cells = list(cells)
        if not any(value is not None and str(value).strip() for value in cells):
            continue
        if len(result) >= limit or line > limit + 100:
            raise ValidationError({label: f"At most {limit} data rows are allowed."})
        if any(str(value or "").strip() for value in cells[len(names):]):
            raise ValidationError({label: f"Row {line}: too many columns."})
        row = dict(zip(names, ["" if value is None else str(value) for value in cells]))
        if any(len(value) > 4096 for value in row.values()):
            raise ValidationError({label: f"Row {line}: a cell exceeds the size limit."})
        result.append({"line": line, **{column: row.get(column, "") for column in columns}})
    return result


def read_text_table(value):
    label, columns, limit = "APIs", API_COLUMNS, MAX_APIS
    if len(value.encode("utf-8")) > MAX_BYTES:
        raise ValidationError({label: "Table exceeds 2 MiB."})
    first = value.splitlines()[0] if value.splitlines() else ""
    delimiter = "\t" if "\t" in first else ","
    try:
        return _rows(csv.reader(io.StringIO(value, newline=""), delimiter=delimiter, strict=True), columns, limit, label)
    except csv.Error:
        raise ValidationError({label: "Invalid CSV/TSV syntax; check quoting and line breaks."}) from None


def read_upload(upload):
    raw = upload.read(MAX_BYTES + 1)
    if len(raw) > MAX_BYTES:
        raise ValidationError({"file": "Files must not exceed 2 MiB."})
    name = str(upload.name).lower()
    if name.endswith(".csv"):
        try:
            return read_text_table(raw.decode("utf-8-sig"))
        except UnicodeDecodeError:
            raise ValidationError({"file": "Save CSV as UTF-8, or use XLSX."}) from None
    if not name.endswith(".xlsx"):
        raise ValidationError({"file": "Use .xlsx or UTF-8 .csv."})
    try:
        with ZipFile(io.BytesIO(raw)) as archive:
            items = archive.infolist()
            if len(items) > 128 or sum(item.file_size for item in items) > MAX_EXPANDED:
                raise ValueError()
            if len({item.filename for item in items}) != len(items):
                raise ValueError()
            for item in items:
                if item.flag_bits & 1 or "vbaproject" in item.filename.lower() or "externallinks" in item.filename.lower():
                    raise ValueError()
                if item.filename.endswith((".xml", ".rels")):
                    xml = archive.read(item)
                    if b"<!DOCTYPE" in xml.upper() or b"<!ENTITY" in xml.upper():
                        raise ValueError()
                    if item.filename.startswith("xl/worksheets/") and item.filename.endswith(".xml"):
                        tree = ElementTree.fromstring(xml)
                        for cell in tree.iter():
                            if cell.tag.rsplit("}", 1)[-1] != "c":
                                continue
                            address = re.fullmatch(r"([A-Z]+)([1-9][0-9]*)", cell.attrib.get("r", ""))
                            if not address:
                                raise ValueError()
                            column = 0
                            for letter in address[1]:
                                column = column * 26 + ord(letter) - 64
                            if column > len(API_COLUMNS) or int(address[2]) > MAX_APIS + 100:
                                raise ValueError()
        book = load_workbook(io.BytesIO(raw), read_only=True, data_only=False, keep_links=False)
        try:
            if "APIs" not in book.sheetnames or set(book.sheetnames) - {"APIs", "Guide"}:
                raise ValidationError({"file": "Use only APIs and optional Guide sheets. Model Offers are discovered by the Runtime, not imported."})
            tables = {}
            for label, columns, limit in [("APIs", API_COLUMNS, MAX_APIS)]:
                if label not in book:
                    tables[label] = []
                    continue
                sheet = book[label]
                if (sheet.max_row or 0) > limit + 100 or (sheet.max_column or 0) > len(columns):
                    raise ValueError()
                # Ignore forged dimension hints and bound iteration explicitly.
                sheet.reset_dimensions()
                values = []
                for cells in sheet.iter_rows(max_row=MAX_APIS + 102, max_col=len(API_COLUMNS) + 1):
                    if any(cell.data_type in {"f", "e"} for cell in cells):
                        raise ValidationError({label: "Formulas and Excel error cells are not accepted. Paste values only."})
                    values.append([cell.value for cell in cells])
                tables[label] = _rows(values, columns, limit, label)
            return tables["APIs"]
        finally:
            book.close()
    except ValidationError:
        raise
    except Exception:
        # Parser errors can contain cell values. Never expose them to API responses/logs.
        raise ValidationError({"file": "Invalid or unsafe XLSX file; use a fresh template without macros or external links."}) from None


def template_xlsx():
    book = Workbook()
    book.remove(book.active)
    for title, columns in [("APIs", API_COLUMNS)]:
        sheet = book.create_sheet(title)
        sheet.append(columns)
        sheet.freeze_panes = "A2"
        for cell in sheet[1]:
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = PatternFill("solid", fgColor="18201B")
            sheet.column_dimensions[cell.column_letter].width = min(38, max(22, len(cell.value) + 2))
    guide = book.create_sheet("Guide")
    for line in ["Nexilume Provider API import", "APIs: one API per api_ref; api_ref is its stable account identity.",
                 "Model Offers are always discovered by the Runtime. Do not supply model lists or model pricing.",
                 "Empty API keys preserve credentials only when updating an existing connection.",
                 "Start the Runtime after import to use the existing automatic model discovery flow. Import does not start it.",
                 "Current Organization/Project is selected in Nexus, never supplied by the spreadsheet.",
                 "100 APIs / 2 MiB per file. Do not include formulas, macros or external links.",
                 "No Model Offer, Marketplace publication, Source, Pool or Router changes.",
                 "This file can contain credentials. Store it securely and delete your local copy when no longer needed."]:
        guide.append([line])
    guide.column_dimensions["A"].width = 120
    output = io.BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()


def csv_bytes(columns, rows=()):
    output = io.StringIO(newline="")
    writer = csv.writer(output)
    writer.writerow(columns)
    for row in rows:
        # Reports are opened by spreadsheet software: neutralize formula injection.
        writer.writerow([("'" + str(v)) if str(v).lstrip().startswith(("=", "+", "-", "@")) else str(v) for v in row])
    return output.getvalue().encode("utf-8-sig")
